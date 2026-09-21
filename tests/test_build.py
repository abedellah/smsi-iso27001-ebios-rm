"""Le script de vérification n'est utile que s'il échoue quand il le doit : chaque test casse les données d'une façon."""
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

import build

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def et():
    return build.Etude()


def erreurs(et):
    return build.verifier(et)[0]


def contient(errs, morceau):
    return any(morceau in e for e in errs)


# -- les données livrées sont cohérentes -------------------------------------------------------------
def test_donnees_livrees_sans_erreur(et):
    errs, warns = build.verifier(et)
    assert errs == [] and warns == []


def test_93_mesures_annexe_a_decidees(et):
    assert len(et.annex) == 93
    assert set(et.soa()) == set(et.annex)


def test_42_regles_du_guide(et):
    assert [r["id"] for r in et.socle] == list(range(1, 43))


def test_tableaux_generes_a_jour(et):
    docs = build.tableaux(et, et.soa())
    for nom, texte in docs.items():
        assert (build.OUT_MD / nom).read_text(encoding="utf-8") == texte, f"{nom} périmé : relancer tools/build.py"


def test_bloc_de_synthese_du_readme_a_jour(et):
    chemin, nouveau = build.readme_avec_bloc(et, et.soa())
    assert nouveau == chemin.read_text(encoding="utf-8")


# -- calculs -------------------------------------------------------------------------------------
def test_dangerosite_retrouve_les_valeurs_du_guide(et):
    p = et.pp
    assert et.dangerosite(p["F3"]) == pytest.approx(3.0)
    assert et.dangerosite(p["F3"], apres=True) == pytest.approx(2.0)
    assert et.dangerosite(p["P3"]) == pytest.approx(2.25)
    assert et.dangerosite(p["F2"], apres=True) == pytest.approx(1.33, abs=0.01)


def test_vraisemblance_recalculee_egale_celle_du_guide(et):
    for rid, attendu in {"R1": 3, "R2": 2, "R3": 4, "R4": 2, "R5": 1}.items():
        assert et.vraisemblance_scenario(rid) == attendu


def test_niveaux_de_risque(et):
    assert [et.niveau(s) for s in (1, 3, 4, 7, 8, 16)] == ["Faible", "Faible", "Moyen", "Moyen", "Élevé", "Élevé"]


# -- cohérence des documents avec les données --------------------------------------------------------------
def test_tableau_des_risques_du_document_06_correspond_aux_donnees(et):
    texte = (ROOT / "docs" / "06-ebios-atelier5-traitement-du-risque.md").read_text(encoding="utf-8")
    for r in et.e["risques"]:
        ligne = next(l for l in texte.splitlines() if l.startswith(f"| {r['id']} |"))
        s0 = r["gravite"] * r["vraisemblance"]
        s1 = r["residuel"]["gravite"] * r["residuel"]["vraisemblance"]
        assert f"{s0} {et.niveau(s0)}" in ligne, r["id"]
        assert f"{s1} " in ligne and et.niveau(s1) in ligne, r["id"]


def test_bilan_du_socle_du_document_02_correspond_aux_donnees(et):
    texte = (ROOT / "docs" / "02-ebios-atelier1-cadrage.md").read_text(encoding="utf-8")
    vert = sum(1 for r in et.socle if r["etat"] == "vert")
    orange = sum(1 for r in et.socle if r["etat"] == "orange")
    rouge = sum(1 for r in et.socle if r["etat"] == "rouge")
    assert (vert, orange, rouge) == (1, 19, 22)
    assert f"{orange} le\nsont avec restrictions et {rouge} ne le sont pas" in texte


# -- la vérification échoue quand elle le doit --------------------------------------------------------------
def test_mesure_annexe_a_sans_decision(et):
    del et.soa_manuel["decisions"]["5.5"]
    assert contient(erreurs(et), "5.5 : aucune décision")


def test_exclusion_sans_approbation(et):
    del et.soa_manuel["exclusions"]["8.28"]["approuve_par"]
    assert contient(erreurs(et), "8.28 : exclusion sans justification ou sans approbation")


def test_mesure_exclue_mais_utilisee_par_le_plan(et):
    et.soa_manuel["exclusions"]["8.13"] = {"pourquoi": "test", "approuve_par": "Direction générale"}
    assert contient(erreurs(et), "8.13 : exclue mais utilisée")


def test_residuel_superieur_a_l_initial(et):
    et.risques["R1"]["residuel"] = {"gravite": 4, "vraisemblance": 4}
    assert contient(erreurs(et), "R1 : le risque résiduel dépasse le risque initial")


def test_residuel_eleve_sans_acceptation(et):
    del et.risques["R3"]["acceptation"]
    assert contient(erreurs(et), "R3 : risque résiduel élevé sans acceptation formelle")


def test_acceptation_sans_date_de_revue(et):
    del et.risques["R3"]["acceptation"]["revue"]
    assert contient(erreurs(et), "R3 : acceptation sans revue")


def test_mesure_sans_risque_ni_evenement(et):
    et.mes["N12"]["er"] = []
    assert contient(erreurs(et), "N12 : la mesure doit traiter au moins un risque")


def test_risque_sans_mesure(et):
    for m in et.mesures:
        if "R5" in m["risques"]:
            m["risques"].remove("R5")
    assert contient(erreurs(et), "R5 : aucune mesure ne traite ce risque")


def test_ecart_du_socle_sans_mesure(et):
    for m in et.mesures:
        if 15 in m["regles"]:
            m["regles"].remove(15)
    assert contient(erreurs(et), "règle 15 : écart non traité")


def test_vraisemblance_declaree_differente_de_celle_calculee(et):
    et.risques["R3"]["vraisemblance"] = 2
    assert contient(erreurs(et), "R3 : vraisemblance 2 déclarée, 4 calculée")


def test_dangerosite_differente_de_celle_du_guide(et):
    et.pp["F3"]["maturite"] = 4
    assert contient(erreurs(et), "F3 : dangerosité")


def test_gravite_du_risque_differente_du_scenario_strategique(et):
    et.risques["R1"]["gravite"] = 4
    assert contient(erreurs(et), "R1 : gravité 4 différente du scénario stratégique")


def test_couple_retenu_sans_scenario_strategique(et):
    et.e["scenarios_strategiques"] = [s for s in et.e["scenarios_strategiques"] if s["id"] != "SS3"]
    assert contient(erreurs(et), "SO5 : couple retenu sans scénario stratégique")


def test_indicateur_de_depart_incoherent(et):
    et.indicateurs[0]["depart"] = "12 sur 42"
    assert contient(erreurs(et), "indicateur I01")


def test_regle_avec_etat_inconnu(et):
    et.regles[10]["etat"] = "bleu"
    assert contient(erreurs(et), "règle 10 : état inconnu")


def test_mesure_avec_mesure_annexe_a_inconnue(et):
    et.mes["N01"]["iso"].append("9.9")
    assert contient(erreurs(et), "N01 : mesure Annexe A inconnue 9.9")


def test_statut_calcule_suit_l_etat_du_plan(et):
    avant = et.soa()["7.10"]["statut"]
    for m in et.mesures:
        m["statut"] = "Terminé"
    for r in et.socle:
        r["etat"] = "vert"
    assert avant != "implemented"
    assert et.soa()["7.10"]["statut"] == "implemented"


# -- classeur Excel ---------------------------------------------------------------------------------
def test_classeur_excel_contient_des_formules(et, tmp_path, monkeypatch):
    sortie = tmp_path / "classeur.xlsx"
    monkeypatch.setattr(build, "OUT_XLSX", sortie)
    build.excel(et, et.soa())
    wb = load_workbook(sortie)
    assert {"Socle", "Risques", "Parties prenantes", "Plan de traitement", "Déclaration d'applicabilité", "Indicateurs"} <= set(wb.sheetnames)
    ws = wb["Risques"]
    assert ws["G2"].value == "=E2*F2"
    assert ws["H2"].value == '=IF(G2>=8,"Élevé",IF(G2>=4,"Moyen","Faible"))'
    assert wb["Parties prenantes"]["H2"].value.startswith("=ROUND(")
    assert wb["Déclaration d'applicabilité"].max_row == 94


# -- les liens entre documents ne sont pas cassés -------------------------------------------------------
def test_liens_markdown_relatifs_valides():
    liens = re.compile(r"\]\(([^)#\s]+)(?:#[^)]*)?\)")
    casses = []
    for md in list(ROOT.glob("*.md")) + list((ROOT / "docs").rglob("*.md")):
        for cible in liens.findall(md.read_text(encoding="utf-8")):
            if cible.startswith(("http://", "https://", "mailto:")):
                continue
            if not (md.parent / cible).resolve().exists():
                casses.append(f"{md.relative_to(ROOT)} -> {cible}")
    assert casses == []


def test_chaque_bloc_mermaid_est_bien_forme():
    """Contrôle simple : blocs ouverts et fermés, sous-graphes fermés. Le rendu lui-même n'est pas testé."""
    for md in list((ROOT / "docs").rglob("*.md")):
        texte = md.read_text(encoding="utf-8")
        assert texte.count("```") % 2 == 0, f"{md.name} : bloc de code non fermé"
        for bloc in re.findall(r"```mermaid\n(.*?)```", texte, re.S):
            assert bloc.count("subgraph ") == len(re.findall(r"^\s*end\s*$", bloc, re.M)), md.name


def test_classeur_non_reecrit_si_contenu_identique(et, tmp_path, monkeypatch):
    sortie = tmp_path / "classeur.xlsx"
    monkeypatch.setattr(build, "OUT_XLSX", sortie)
    build.excel(et, et.soa())
    avant = sortie.read_bytes()
    build.excel(et, et.soa())
    assert sortie.read_bytes() == avant
