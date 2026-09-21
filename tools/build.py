"""Vérifie la cohérence des données puis génère les tableaux (docs/tableaux) et le classeur Excel.

Usage :  python tools/build.py            vérifie et génère
         python tools/build.py --check    vérifie, et échoue si les tableaux générés sont périmés
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter, defaultdict
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
OUT_MD = ROOT / "docs" / "tableaux"
OUT_XLSX = ROOT / "reports" / "SMSI-biotech-classeur.xlsx"

ETATS = ("vert", "orange", "rouge")
STATUTS_MESURE = ("Terminé", "En cours", "À lancer")
POIDS_ETAT = {"vert": 1.0, "orange": 0.5, "rouge": 0.0}
POIDS_MESURE = {"Terminé": 1.0, "En cours": 0.5, "À lancer": 0.0}
LIBELLE_STATUT = {"implemented": "Mise en œuvre", "partial": "Partielle", "planned": "À mettre en œuvre",
                  "excluded": "Exclue"}
BASES = {"risque", "socle", "légal", "contractuel", "bonne pratique"}


def read(name):
    with (DATA / name).open(encoding="utf-8") as fh:
        return yaml.safe_load(fh)


class Etude:
    def __init__(self):
        self.socle = read("socle.yaml")["rules"]
        self.e = read("ebios.yaml")
        self.mesures = read("mesures.yaml")["mesures"]
        self.soa_manuel = read("soa_manuel.yaml")
        self.indicateurs = read("indicateurs.yaml")["indicateurs"]
        annex = read("annex_a.yaml")
        self.annex = {}
        for theme, controls in annex.items():
            for cid, titre in controls.items():
                self.annex[str(cid)] = {"titre": titre, "theme": theme}
        self.regles = {r["id"]: r for r in self.socle}
        self.risques = {r["id"]: r for r in self.e["risques"]}
        self.ers = {r["id"]: r for r in self.e["evenements_redoutes"]}
        self.mes = {m["id"]: m for m in self.mesures}
        self.pp = {p["id"]: p for p in self.e["parties_prenantes"]}
        self.so = {s["id"]: s for s in self.e["couples_sr_ov"]}
        self.chemins = {c["id"]: (ss, c) for ss in self.e["scenarios_strategiques"] for c in ss["chemins"]}

    # -- calculs ------------------------------------------------------------------------------
    def niveau(self, score):
        score = round(score)
        for n in self.e["niveaux_risque"]:
            if n["min"] <= score <= n["max"]:
                return n["nom"]
        return "?"

    @staticmethod
    def score(g, v):
        return g * v

    def dangerosite(self, p, apres=False):
        c = dict(dependance=p["dependance"], penetration=p["penetration"], maturite=p["maturite"],
                 confiance=p["confiance"])
        if apres:
            c.update(p.get("apres", {}))
        return c["dependance"] * c["penetration"] / (c["maturite"] * c["confiance"])

    def zone(self, d):
        z = self.e["zones"]
        return "Danger" if d >= z["danger"] else "Contrôle" if d >= z["controle"] else "Veille" if d >= z["veille"] else "Hors seuil"

    def vraisemblance_mode(self, mode):
        return min(a["v"] for a in mode["actions"])

    def vraisemblance_scenario(self, rid):
        return max(self.vraisemblance_mode(m) for m in self.e["operationnels"][rid]["modes"])

    def mesures_du_risque(self, rid):
        return [m for m in self.mesures if rid in m.get("risques", [])]

    def controles_lies(self):
        """Pour chaque mesure de l'Annexe A : règles du socle, mesures et risques qui la justifient."""
        lies = {cid: {"regles": set(), "mesures": set(), "risques": set()} for cid in self.annex}
        for r in self.socle:
            for cid in r["iso"]:
                if cid in lies:
                    lies[cid]["regles"].add(r["id"])
        for m in self.mesures:
            for cid in m["iso"]:
                if cid in lies:
                    lies[cid]["mesures"].add(m["id"])
                    lies[cid]["risques"].update(m.get("risques", []))
        return lies

    def soa(self):
        """Décision, base et statut de chacune des 93 mesures de l'Annexe A."""
        lies = self.controles_lies()
        manuel = self.soa_manuel.get("decisions", {})
        exclus = self.soa_manuel.get("exclusions", {})
        out = {}
        for cid, a in self.annex.items():
            l = lies[cid]
            if cid in exclus:
                out[cid] = {"statut": "excluded", "base": [], "pourquoi": exclus[cid]["pourquoi"],
                            "approuve_par": exclus[cid]["approuve_par"], **l}
                continue
            items = [POIDS_ETAT[self.regles[r]["etat"]] for r in l["regles"]]
            items += [POIDS_MESURE[self.mes[m]["statut"]] for m in l["mesures"]]
            base, morceaux = [], []
            man = manuel.get(cid, {})
            if l["risques"]:
                base.append("risque")
                morceaux.append("Retenue pour traiter " + ", ".join(sorted(l["risques"])))
            if l["regles"]:
                base.append("socle")
                morceaux.append(("corrige l'écart à la règle " if len(l["regles"]) == 1 else "corrige l'écart aux règles ")
                                + ", ".join(str(r) for r in sorted(l["regles"])) + " du Guide d'hygiène")
            for b in man.get("base", []):
                if b not in base:
                    base.append(b)
            texte = " ; ".join(morceaux)
            if texte:
                texte = texte[0].upper() + texte[1:] + "."
            if man.get("pourquoi"):
                texte = (texte + " " if texte else "") + man["pourquoi"]
            if items:
                moy = sum(items) / len(items)
                statut = "implemented" if moy >= 0.999 else "partial" if moy > 0 else "planned"
            else:
                statut = man.get("statut", "planned")
            if "statut" in man:
                statut = man["statut"]
            out[cid] = {"statut": statut, "base": base, "pourquoi": texte, "approuve_par": "", **l}
        return out


# -- vérifications --------------------------------------------------------------------------------
def verifier(et: Etude):
    errs, warns = [], []
    err, warn = errs.append, warns.append
    annex = et.annex
    e = et.e

    # règles du socle
    if [r["id"] for r in et.socle] != list(range(1, 43)):
        err("socle : les 42 règles du guide doivent être présentes, numérotées de 1 à 42")
    for r in et.socle:
        if r["etat"] not in ETATS:
            err(f"règle {r['id']} : état inconnu {r['etat']!r}")
        if r["etat"] != "vert" and not r.get("ecart"):
            err(f"règle {r['id']} : un écart doit être décrit")
        for cid in r["iso"]:
            if cid not in annex:
                err(f"règle {r['id']} : mesure Annexe A inconnue {cid}")

    # mesures
    for m in et.mesures:
        mid = m["id"]
        if m["statut"] not in STATUTS_MESURE:
            err(f"{mid} : statut inconnu {m['statut']!r}")
        if not m.get("risques") and not m.get("er"):
            err(f"{mid} : la mesure doit traiter au moins un risque ou un événement redouté")
        for rid in m.get("risques", []):
            if rid not in et.risques:
                err(f"{mid} : risque inconnu {rid}")
        for er in m.get("er", []):
            if er not in et.ers:
                err(f"{mid} : événement redouté inconnu {er}")
        for cid in m["iso"]:
            if cid not in annex:
                err(f"{mid} : mesure Annexe A inconnue {cid}")
        if not m["iso"]:
            err(f"{mid} : aucune mesure de l'Annexe A associée")
        for r in m["regles"]:
            if r not in et.regles:
                err(f"{mid} : règle du socle inconnue {r}")
        if m["cout"] not in ("+", "++", "+++", "++++"):
            err(f"{mid} : coût inconnu {m['cout']!r}")
        if m["priorite"] not in ("P1", "P2", "P3"):
            err(f"{mid} : priorité inconnue")
        if not isinstance(m.get("echeance_mois"), int):
            err(f"{mid} : échéance manquante")
    # chaque écart du socle est traité ou explicitement accepté
    couvertes = {r for m in et.mesures for r in m["regles"]}
    for r in et.socle:
        if r["etat"] != "vert" and r["id"] not in couvertes and not r.get("accepte"):
            err(f"règle {r['id']} : écart non traité par une mesure et non accepté")

    # événements redoutés et couples SR/OV
    for er in et.ers.values():
        if er["vm"] not in {v["id"] for v in e["valeurs_metier"]}:
            err(f"{er['id']} : valeur métier inconnue")
    for so in e["couples_sr_ov"]:
        if so["er"] not in et.ers:
            err(f"{so['id']} : événement redouté inconnu")
    ss_par_so = {ss["so"]: ss for ss in e["scenarios_strategiques"]}
    for so in e["couples_sr_ov"]:
        if so["retenu"] and so["id"] not in ss_par_so:
            err(f"{so['id']} : couple retenu sans scénario stratégique")
    for ss in e["scenarios_strategiques"]:
        if ss["gravite"] != et.ers[et.so[ss["so"]]["er"]]["gravite"]:
            err(f"{ss['id']} : gravité différente de celle de l'événement redouté associé")
        for c in ss["chemins"]:
            for pid in c["pp"]:
                if not et.pp[pid]["critique"]:
                    err(f"{c['id']} : {pid} n'est pas une partie prenante critique")

    # parties prenantes : les 3 niveaux de menace du guide (p. 54)
    attendu = {"F2": (2.0, 1.33), "F3": (3.0, 2.0), "P3": (2.25, 1.5)}
    for pid, (init, res) in attendu.items():
        p = et.pp[pid]
        if abs(et.dangerosite(p) - init) > 0.01 or abs(et.dangerosite(p, True) - res) > 0.01:
            err(f"{pid} : dangerosité {et.dangerosite(p):.2f}/{et.dangerosite(p, True):.2f} "
                f"différente de celle du guide {init}/{res}")

    # risques
    for r in e["risques"]:
        rid = r["id"]
        if r["chemin"] not in et.chemins:
            err(f"{rid} : chemin d'attaque inconnu")
        else:
            ss, _ = et.chemins[r["chemin"]]
            if r["gravite"] != ss["gravite"]:
                err(f"{rid} : gravité {r['gravite']} différente du scénario stratégique {ss['gravite']}")
        if r["er"] not in et.ers:
            err(f"{rid} : événement redouté inconnu")
        for k in ("gravite", "vraisemblance"):
            if not 1 <= r[k] <= 4:
                err(f"{rid} : {k} hors échelle")
        res = r["residuel"]
        if not (1 <= res["gravite"] <= 4 and 1 <= res["vraisemblance"] <= 4):
            err(f"{rid} : résiduel hors échelle")
        if et.score(res["gravite"], res["vraisemblance"]) > et.score(r["gravite"], r["vraisemblance"]):
            err(f"{rid} : le risque résiduel dépasse le risque initial")
        if not et.mesures_du_risque(rid):
            err(f"{rid} : aucune mesure ne traite ce risque")
        niveau_res = et.niveau(et.score(res["gravite"], res["vraisemblance"]))
        if niveau_res == "Élevé":
            a = r.get("acceptation")
            if not a:
                err(f"{rid} : risque résiduel élevé sans acceptation formelle de la direction")
            else:
                for k in ("par", "date", "revue", "decision_prevue"):
                    if not a.get(k):
                        err(f"{rid} : acceptation sans {k}")
        elif r.get("acceptation"):
            warn(f"{rid} : acceptation inutile, le niveau résiduel est {niveau_res}")
        # vraisemblance issue des modes opératoires
        if rid not in e["operationnels"]:
            err(f"{rid} : pas de scénario opérationnel")
        else:
            calc = et.vraisemblance_scenario(rid)
            if calc != r["vraisemblance"]:
                err(f"{rid} : vraisemblance {r['vraisemblance']} déclarée, {calc} calculée à partir des modes opératoires")
            for m in e["operationnels"][rid]["modes"]:
                for a in m["actions"]:
                    for rg in a["socle"]:
                        if rg not in et.regles:
                            err(f"{rid} : règle du socle inconnue {rg} dans un mode opératoire")

    # déclaration d'applicabilité
    manuel = et.soa_manuel.get("decisions", {})
    exclus = et.soa_manuel.get("exclusions", {})
    lies = et.controles_lies()
    for cid in list(manuel) + list(exclus):
        if cid not in annex:
            err(f"déclaration d'applicabilité : {cid} n'est pas une mesure de l'Annexe A")
    for cid, l in lies.items():
        li = l["regles"] or l["mesures"]
        if cid in exclus:
            if l["mesures"] or l["risques"]:
                err(f"{cid} : exclue mais utilisée par {sorted(l['mesures'])}")
            if not exclus[cid].get("pourquoi") or not exclus[cid].get("approuve_par"):
                err(f"{cid} : exclusion sans justification ou sans approbation")
        elif not li and cid not in manuel:
            err(f"{cid} : aucune décision (ni rattachée à une règle ou une mesure, ni décrite dans soa_manuel.yaml)")
    for cid, d in manuel.items():
        for b in d.get("base", []):
            if b not in BASES:
                err(f"{cid} : base inconnue {b!r}")
        if cid not in exclus and not lies.get(cid, {}).get("mesures") and not lies.get(cid, {}).get("regles"):
            if d.get("statut") not in ("implemented", "partial", "planned"):
                err(f"{cid} : statut manuel requis")
            if not d.get("pourquoi") or not d.get("base"):
                err(f"{cid} : justification et base requises")

    # indicateurs : les valeurs de départ chiffrées doivent correspondre aux données
    depart = {i["id"]: i.get("depart", "") for i in et.indicateurs}
    vert = sum(1 for r in et.socle if r["etat"] == "vert")
    termine = sum(1 for m in et.mesures if m["statut"] == "Terminé")
    eleves = sum(1 for r in e["risques"] if et.niveau(r["gravite"] * r["vraisemblance"]) == "Élevé")
    attendus = {"I01": f"{vert} sur {len(et.socle)}", "I02": f"{termine} sur {len(et.mesures)}", "I03": f"{eleves} avant traitement"}
    for iid, val in attendus.items():
        if depart.get(iid) != val:
            err(f"indicateur {iid} : départ {depart.get(iid)!r}, les données donnent {val!r}")
    # indicateurs
    for i in et.indicateurs:
        for k in ("id", "nom", "formule", "cible", "frequence", "responsable"):
            if not i.get(k):
                err(f"indicateur {i.get('id')} : {k} manquant")
    return errs, warns


# -- génération des tableaux --------------------------------------------------------------------------
def md_table(head, rows):
    out = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    for r in rows:
        out.append("| " + " | ".join(str(c).replace("|", "/").replace("\n", " ") for c in r) + " |")
    return "\n".join(out)


ENTETE = "*Fichier généré par `python tools/build.py` à partir de `data/`. Ne pas modifier à la main.*"


def grille(et, pick):
    cell = defaultdict(list)
    for r in et.e["risques"]:
        g, v = pick(r)
        cell[(g, v)].append(r["id"])
    lignes = ["| Gravité \\ Vraisemblance | V1 | V2 | V3 | V4 |", "|---|---|---|---|---|"]
    for g in (4, 3, 2, 1):
        lignes.append(f"| **G{g}** | " + " | ".join(", ".join(cell.get((g, v), [])) or "." for v in (1, 2, 3, 4)) + " |")
    return "\n".join(lignes)


def tableaux(et: Etude, soa):
    docs = {}
    # socle
    c = Counter(r["etat"] for r in et.socle)
    lignes = [
        "# Socle de sécurité : état d'application du Guide d'hygiène de l'ANSSI", "", ENTETE, "",
        "Référentiel : ANSSI, *Guide d'hygiène informatique*, v2.0, septembre 2017 (Licence Ouverte Etalab V1). "
        "Niveau évalué : standard. Les états sont ceux de l'entreprise fictive de biotechnologie ; ils s'appuient sur le "
        "guide EBIOS RM quand celui-ci en dit quelque chose, et sont sinon une hypothèse cohérente avec sa maturité faible "
        "(voir la colonne Source).", "",
        f"**Bilan : {c['vert']} règle appliquée sans restriction, {c['orange']} avec restrictions, {c['rouge']} non appliquées "
        f"(sur {len(et.socle)}).**", ""]
    traitees = defaultdict(list)
    for m in et.mesures:
        for r in m["regles"]:
            traitees[r].append(m["id"])
    section = None
    rows = []
    for r in et.socle:
        if r["section"] != section:
            if rows:
                lignes += [md_table(["Règle", "État", "Écart", "Mesures", "Annexe A ISO 27001", "Source"], rows), ""]
                rows = []
            section = r["section"]
            lignes.append(f"## {section}")
            lignes.append("")
        rows.append([f"{r['id']}. {r['titre']}", r["etat"], r["ecart"] or "-", ", ".join(traitees.get(r["id"], [])) or "-",
                     ", ".join(r["iso"]) or "clause 6.1.2", r["source"]])
    lignes += [md_table(["Règle", "État", "Écart", "Mesures", "Annexe A ISO 27001", "Source"], rows), ""]
    docs["socle.md"] = "\n".join(lignes)

    # parties prenantes
    rows = []
    for p in et.e["parties_prenantes"]:
        d0, d1 = et.dangerosite(p), et.dangerosite(p, True)
        rows.append([p["id"], p["nom"], p["categorie"], p["dependance"], p["penetration"], p["maturite"], p["confiance"],
                     f"{d0:.2f}".replace(".", ","), et.zone(d0), "oui" if p["critique"] else "non",
                     f"{d1:.2f}".replace(".", ",") if p.get("apres") else "-", p["source"]])
    docs["parties-prenantes.md"] = "\n".join([
        "# Cartographie de dangerosité de l'écosystème", "", ENTETE, "",
        "Dangerosité = (dépendance x pénétration) / (maturité cyber x confiance), chaque composante de 1 à 4. "
        "Seuils du guide : veille 0,2 ; contrôle 0,9 ; danger 2,5. Les composantes de F2, F3 et P3 sont reconstituées pour "
        "retrouver les niveaux donnés par le guide (2, 3 et 2,25 puis 1,3, 2 et 1,5) ; les autres sont des hypothèses.", "",
        md_table(["Réf.", "Partie prenante", "Catégorie", "Dép.", "Pén.", "Mat.", "Conf.", "Dangerosité", "Zone", "Critique",
                  "Après mesures", "Source"], rows), ""])

    # registre des risques
    rows = []
    for r in et.e["risques"]:
        g, v, rg, rv = r["gravite"], r["vraisemblance"], r["residuel"]["gravite"], r["residuel"]["vraisemblance"]
        s0, s1 = g * v, rg * rv
        rows.append([r["id"], r["libelle"], f"G{g}", f"V{v}", f"{s0} {et.niveau(s0)}", f"G{rg}", f"V{rv}", f"{s1} {et.niveau(s1)}",
                     ", ".join(m["id"] for m in et.mesures_du_risque(r["id"])), r["source"]])
    lg = ["# Registre des risques", "", ENTETE, "",
          "Score = gravité x vraisemblance. " + " ; ".join(f"{n['nom']} {n['min']} à {n['max']} ({n['decision']})" for n in et.e["niveaux_risque"]) + ".", "",
          md_table(["Réf.", "Scénario de risque", "G", "V", "Niveau initial", "G rés.", "V rés.", "Niveau résiduel", "Mesures", "Source"], rows), "",
          "## Cartographie du risque initial", "", grille(et, lambda r: (r["gravite"], r["vraisemblance"])), "",
          "## Cartographie du risque résiduel", "", grille(et, lambda r: (r["residuel"]["gravite"], r["residuel"]["vraisemblance"])), ""]
    docs["registre-risques.md"] = "\n".join(lg)

    # modes opératoires : vraisemblances calculées
    rows = []
    for rid, blob in et.e["operationnels"].items():
        for i, m in enumerate(blob["modes"], 1):
            rows.append([rid, f"Mode {i}", m["nom"], " > ".join(f"V{a['v']}" for a in m["actions"]), f"V{et.vraisemblance_mode(m)}"])
        rows.append([rid, "**Scénario**", "Mode de moindre effort", "", f"**V{et.vraisemblance_scenario(rid)}**"])
    docs["vraisemblances.md"] = "\n".join([
        "# Vraisemblance des scénarios opérationnels", "", ENTETE, "",
        "Vraisemblance d'un mode = la plus faible de ses actions ; vraisemblance du scénario = celle du mode de moindre "
        "effort pour l'attaquant (guide EBIOS RM, p. 65).", "",
        md_table(["Risque", "Mode", "Description", "Actions élémentaires", "Vraisemblance"], rows), ""])

    # plan de traitement
    rows = []
    for m in et.mesures:
        rows.append([m["id"], m["titre"], m["groupe"], ", ".join(m["risques"] or m.get("er", [])), m["responsable"],
                     m["freins"] or "-", m["cout"], m["charge"] or "-", f"T0+{m['echeance_mois']} mois", m["priorite"],
                     m.get("priorite_anssi") or "-", m["statut"], ", ".join(str(r) for r in m["regles"]) or "-",
                     ", ".join(m["iso"]), m["source"]])
    docs["plan-traitement.md"] = "\n".join([
        "# Plan de traitement du risque", "", ENTETE, "",
        "M01 à M13 : mesures du guide EBIOS RM (pp. 54 et 77), reprises avec leurs statuts. N01 à N13 : mesures ajoutées par "
        "ce projet. Coût : + moins de 10 k€, ++ de 10 à 50 k€, +++ de 50 à 200 k€ (échelle de ce projet, hypothèse). T0 = 1er octobre 2026.", "",
        md_table(["Réf.", "Mesure", "Axe", "Risques", "Responsable", "Freins", "Coût", "Charge", "Échéance", "Priorité",
                  "Priorité guide", "Statut", "Règles du socle", "Annexe A", "Source"], rows), ""])

    # déclaration d'applicabilité
    rows = []
    for cid, a in et.annex.items():
        d = soa[cid]
        rows.append([cid, a["titre"], a["theme"], "Non (exclue)" if d["statut"] == "excluded" else "Oui", LIBELLE_STATUT[d["statut"]],
                     ", ".join(d["base"]) or "-", d["pourquoi"] or "-", ", ".join(sorted(d["mesures"])) or "-",
                     ", ".join(sorted(d["risques"])) or "-", d["approuve_par"] or "-"])
    cs = Counter(d["statut"] for d in soa.values())
    docs["declaration-applicabilite.md"] = "\n".join([
        "# Déclaration d'applicabilité (ISO/IEC 27001:2022, Annexe A)", "", ENTETE, "",
        "Les intitulés sont des paraphrases courtes de l'Annexe A, pas le texte de la norme. Le statut est calculé à partir de "
        "l'état des règles du socle et des mesures du plan de traitement liées à la mesure de l'Annexe A ; les mesures sans lien "
        "sont décrites dans `data/soa_manuel.yaml`.", "",
        f"**Bilan : {cs['implemented']} mises en œuvre, {cs['partial']} partielles, {cs['planned']} à mettre en œuvre, "
        f"{cs['excluded']} exclues (sur {len(soa)}).**", "",
        md_table(["N°", "Mesure", "Thème", "Applicable", "Statut", "Base du choix", "Justification", "Mesures du plan", "Risques", "Exclusion approuvée par"], rows), ""])
    docs["scenarios-operationnels.md"] = mermaid(et)
    docs["trajectoire-et-budget.md"] = trajectoire(et)
    docs["risques-residuels.md"] = residuels(et)
    docs["indicateurs.md"] = indicateurs_md(et)
    return docs


# -- contenus dérivés : graphes d'attaque, résiduels, trajectoire, budget, synthèse --------------------
COUT_BORNES = {"+": (0, 10), "++": (10, 50), "+++": (50, 200)}


MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]


def date_fr(d):
    return f"{d.day}{'er' if d.day == 1 else ''} {MOIS[d.month - 1]} {d.year}"


def nettoyer(t):
    return str(t).replace('"', "'")


def mermaid(et):
    out = ["# Scénarios opérationnels : graphes d'attaque", "", ENTETE, "",
           "Un graphe par risque. Chaque bloc est un mode opératoire, lu de gauche à droite : de la reconnaissance jusqu'à l'exploitation. "
           "V1 à V4 est la vraisemblance élémentaire de l'action ; « règles » renvoie aux règles du Guide d'hygiène dont l'écart la facilite "
           "(voir [le socle](socle.md)). La vraisemblance du scénario est celle du mode de moindre effort, calculée "
           "dans [vraisemblances.md](vraisemblances.md).", ""]
    for rid, blob in et.e["operationnels"].items():
        r = et.risques[rid]
        out += [f"## {rid} : {r['libelle']}", "",
                f"Gravité G{r['gravite']}, vraisemblance V{r['vraisemblance']} ({r['source']}).", "", "```mermaid", "flowchart LR"]
        for mi, m in enumerate(blob["modes"], 1):
            out.append(f'  subgraph {rid}_M{mi}["Mode {mi} : {nettoyer(m["nom"])}"]')
            out.append("    direction LR")
            for ai, a in enumerate(m["actions"], 1):
                regles = f"<br/>règles {', '.join(str(x) for x in a['socle'])}" if a["socle"] else ""
                out.append(f'    {rid}_{mi}_{ai}["{nettoyer(a["t"])}<br/>V{a["v"]}{regles}"]')
            out.append("    " + " --> ".join(f"{rid}_{mi}_{ai}" for ai in range(1, len(m["actions"]) + 1)))
            out.append("  end")
        out += ["```", ""]
    return "\n".join(out)


def trajectoire(et):
    mois = (0, 6, 12, 18)
    lignes = []
    for r in et.e["risques"]:
        ms = et.mesures_du_risque(r["id"])
        s0 = r["gravite"] * r["vraisemblance"]
        s1 = r["residuel"]["gravite"] * r["residuel"]["vraisemblance"]
        vals = []
        for m in mois:
            fait = sum(1 for x in ms if x["statut"] == "Terminé" or (m > 0 and x["echeance_mois"] <= m))
            f = fait / len(ms)
            vals.append(round(s0 - (s0 - s1) * f, 1))
        lignes.append((r["id"], s0, vals, s1))
    fr = lambda x: f"{x:.1f}".replace(".", ",")
    rows = []
    for rid, s0, vals, s1 in lignes:
        rows.append([rid, s0] + [f"{fr(v)} {et.niveau(v)}" for v in vals] + [f"{s1} {et.niveau(s1)}"])
    moy = [sum(l[2][i] for l in lignes) / len(lignes) for i in range(len(mois))]
    eleves = [sum(1 for l in lignes if round(l[2][i]) >= 8) for i in range(len(mois))]
    rows.append(["**Moyenne**", fr(sum(l[1] for l in lignes) / len(lignes))] + [fr(v) for v in moy] + [fr(sum(l[3] for l in lignes) / len(lignes))])
    rows.append(["**Risques au niveau Élevé**", sum(1 for l in lignes if l[1] >= 8)] + eleves + [sum(1 for l in lignes if l[3] >= 8)])
    reste = [m for m in et.mesures if m["statut"] != "Terminé"]
    bas = sum(COUT_BORNES[m["cout"]][0] for m in reste)
    haut = sum(COUT_BORNES[m["cout"]][1] for m in reste)
    par_prio = defaultdict(lambda: [0, 0, 0])
    for m in reste:
        par_prio[m["priorite"]][0] += 1
        par_prio[m["priorite"]][1] += COUT_BORNES[m["cout"]][0]
        par_prio[m["priorite"]][2] += COUT_BORNES[m["cout"]][1]
    jh = sum(int(m["charge"].split()[0]) for m in reste if m["charge"])
    prio_rows = [[p, v[0], f"{v[1]} à {v[2]} k€"] for p, v in sorted(par_prio.items())]
    return "\n".join([
        "# Trajectoire du risque et enveloppe budgétaire", "", ENTETE, "",
        "Le guide EBIOS RM recommande d'associer une cartographie des risques résiduels à chaque grand jalon du plan (p. 79). "
        "Modèle simple de ce projet (hypothèse) : à chaque jalon, un risque est réduit de son niveau initial vers son niveau cible "
        "en proportion des mesures qui le traitent et qui sont terminées ou dont l'échéance est atteinte. Une mesure n'a d'effet qu'une "
        "fois mise en œuvre (guide, p. 53, note 26). C'est un outil de lecture, pas une mesure du risque réel.", "",
        md_table(["Risque", "Score initial", "T0 (aujourd'hui)", "T0+6 mois", "T0+12 mois", "T0+18 mois", "Score cible"], rows), "",
        "## Enveloppe budgétaire estimée", "",
        f"Mesures restant à réaliser ({len(reste)} sur {len(et.mesures)}) : entre **{bas} et {haut} k€**, et {jh} jours-hommes chiffrés "
        "(le guide ne donne des charges que pour deux mesures ; les autres sont des estimations de ce projet). "
        "Les bornes viennent de l'échelle + / ++ / +++ définie dans le plan (hypothèse).", "",
        md_table(["Priorité", "Mesures", "Budget"], prio_rows), ""])


def residuels(et):
    out = ["# Risques résiduels", "", ENTETE, "",
           "Fiches établies sur le modèle proposé par le guide EBIOS RM (p. 78). Les niveaux résiduels sont des cibles : ils ne "
           "sont atteints que lorsque les mesures listées sont mises en œuvre.", ""]
    for n, r in enumerate(et.e["risques"], 1):
        g0, v0, g1, v1 = r["gravite"], r["vraisemblance"], r["residuel"]["gravite"], r["residuel"]["vraisemblance"]
        s0, s1 = g0 * v0, g1 * v1
        niv = et.niveau(s1)
        ms = et.mesures_du_risque(r["id"])
        out += [f"## RR{n:02d} ({r['id']}) : {r['libelle']}", "",
                f"- **Événement redouté :** {r['er']} : {et.ers[r['er']]['libelle']}",
                f"- **Analyse :** {r['justification_residuel']}.",
                "- **Mesures :** " + ", ".join(m["id"] for m in ms) + ".",
                f"- **Estimation :** gravité G{g0} puis G{g1} ; vraisemblance V{v0} puis V{v1} ; niveau {et.niveau(s0)} ({s0}) puis **{niv}** ({s1}).",
                ]
        a = r.get("acceptation")
        if a:
            out.append(f"- **Gestion du risque résiduel :** accepté par la {a['par']} le {date_fr(a['date'])} ; revue au plus tard le {date_fr(a['revue'])}. "
                       f"Décision à prendre : {a['decision_prevue']}. ({a['source']})")
        elif niv == "Moyen":
            out.append("- **Gestion du risque résiduel :** tolérable sous contrôle ; suivi semestriel au comité de pilotage, réévaluation à chaque cycle opérationnel.")
        else:
            out.append("- **Gestion du risque résiduel :** acceptable en l'état ; réexamen à chaque cycle stratégique.")
        out.append("")
    return "\n".join(out)


def indicateurs_md(et):
    rows = [[i["id"], i["nom"], i["formule"], i.get("depart", "-"), i["cible"], i["frequence"], i["responsable"], i.get("lien", "-")] for i in et.indicateurs]
    return "\n".join(["# Indicateurs de pilotage", "", ENTETE, "",
                      "Cibles : hypothèses de ce projet. Situation de départ : celle de l'entreprise fictive au lancement du plan.", "",
                      md_table(["Réf.", "Indicateur", "Formule", "Départ", "Cible", "Fréquence", "Responsable", "Lien"], rows), ""])


def synthese(et, soa):
    c = Counter(r["etat"] for r in et.socle)
    cs = Counter(d["statut"] for d in soa.values())
    n_orig = lambda lst: sum(1 for x in lst if str(x.get("source", "")).startswith("Original"))
    ini = Counter(et.niveau(r["gravite"] * r["vraisemblance"]) for r in et.e["risques"])
    res = Counter(et.niveau(r["residuel"]["gravite"] * r["residuel"]["vraisemblance"]) for r in et.e["risques"])
    moy0 = sum(r["gravite"] * r["vraisemblance"] for r in et.e["risques"]) / len(et.e["risques"])
    moy1 = sum(r["residuel"]["gravite"] * r["residuel"]["vraisemblance"] for r in et.e["risques"]) / len(et.e["risques"])
    st = Counter(m["statut"] for m in et.mesures)
    ajoutes = [m for m in et.mesures if m["source"].startswith("Original")]
    fr = lambda x: f"{x:.1f}".replace(".", ",")
    return "\n".join([
        f"- **Socle de sécurité** (Guide d'hygiène de l'ANSSI) : {c['vert']} règle sur {len(et.socle)} appliquée sans restriction, "
        f"{c['orange']} avec restrictions, {c['rouge']} non appliquées.",
        f"- **Risques** : {len(et.e['risques'])} scénarios (5 du guide, {n_orig(et.e['risques'])} ajoutés). Avant traitement : "
        f"{ini['Élevé']} de niveau élevé, {ini['Moyen']} moyen ; score moyen {fr(moy0)}. Après traitement : {res['Élevé']} élevé "
        f"(R3, accepté par la direction avec revue datée), {res['Moyen']} moyens, {res['Faible']} faibles ; score moyen {fr(moy1)}.",
        f"- **Plan de traitement** : {len(et.mesures)} mesures, dont {len(et.mesures) - len(ajoutes)} venues du guide et {len(ajoutes)} ajoutées ; "
        f"{st['Terminé']} terminées, {st['En cours']} en cours, {st['À lancer']} à lancer.",
        f"- **ISO/IEC 27001, Annexe A** : {len(soa)} mesures décidées, dont {cs['partial']} partiellement en place, {cs['planned']} à mettre en œuvre "
        f"et {cs['excluded']} exclues avec justification et approbation. Aucune n'est encore pleinement en place : c'est un SMSI à construire.",
        f"- **Ajouts à l'exemple du guide** : {n_orig(et.e['evenements_redoutes'])} événements redoutés, {n_orig(et.e['couples_sr_ov'])} couple source de risque / "
        f"objectif visé, {n_orig(et.e['scenarios_strategiques'])} scénario stratégique, {n_orig(et.e['risques'])} scénarios de risque, "
        f"{len(et.e['operationnels']) - 1} scénarios opérationnels détaillés sur {len(et.e['operationnels'])}, {len(ajoutes)} mesures.",
    ])


BLOC_DEBUT, BLOC_FIN = "<!-- synthese:debut -->", "<!-- synthese:fin -->"


def readme_avec_bloc(et, soa):
    p = ROOT / "README.md"
    t = p.read_text(encoding="utf-8")
    if BLOC_DEBUT not in t or BLOC_FIN not in t:
        return p, t
    avant, reste = t.split(BLOC_DEBUT, 1)
    _, apres = reste.split(BLOC_FIN, 1)
    return p, avant + BLOC_DEBUT + "\n" + synthese(et, soa) + "\n" + BLOC_FIN + apres


# -- Excel ----------------------------------------------------------------------------------------------
def formule_niveau(cell, et):
    bandes = sorted(et.e["niveaux_risque"], key=lambda n: n["min"])
    expr = f'"{bandes[0]["nom"]}"'
    for n in bandes[1:]:  # du plus bas au plus haut : le seuil le plus élevé est testé en premier
        expr = f'IF({cell}>={n["min"]},"{n["nom"]}",{expr})'
    return "=" + expr


def contenu(wb):
    # une chaîne vide est relue comme une cellule vide : on les traite pareil
    return {ws.title: [[None if c.value == "" else c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def ecrire_si_different(wb, chemin):
    """Un classeur Excel n'est jamais identique octet pour octet (dates internes) : on ne le réécrit que si son contenu change,
    pour que Git ne signale pas de modification à chaque exécution."""
    if chemin.exists():
        from openpyxl import load_workbook
        if contenu(load_workbook(chemin)) == contenu(wb):
            return False
    wb.save(chemin)
    return True


def excel(et: Etude, soa):
    wb = Workbook()
    gras, fond = Font(bold=True), PatternFill("solid", fgColor="D9E1F2")

    def feuille(titre, entetes, lignes, largeurs=None):
        ws = wb.create_sheet(titre)
        ws.append(entetes)
        for l in lignes:
            ws.append(l)
        for c in ws[1]:
            c.font, c.fill = gras, fond
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        for i, h in enumerate(entetes, 1):
            ws.column_dimensions[get_column_letter(i)].width = (largeurs or {}).get(get_column_letter(i), max(10, min(len(str(h)) + 2, 24)))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.auto_filter.ref = ws.dimensions
        return ws

    lisez = wb.active
    lisez.title = "Lisez-moi"
    for l in [
        ["SMSI ISO/IEC 27001 et analyse de risques EBIOS RM : cas de référence de l'ANSSI (biotechnologie, fictive)"],
        ["Classeur généré par tools/build.py à partir des fichiers de data/. Les scores et niveaux des feuilles Risques et Parties prenantes sont des formules."],
        ["Source des éléments repris : ANSSI, La méthode EBIOS Risk Manager - Le guide, v1.5, sept. 2024, et Guide d'hygiène informatique, v2.0, sept. 2017. Licence Ouverte Etalab V1. Voir CREDITS.md."],
        ["Colonne Source : 'ANSSI-EBIOS p. N' = repris du guide ; 'Original' = travail de ce projet ; 'Hypothèse' = choix à discuter."],
        ["Projet non affilié à l'ANSSI, au Club EBIOS ni à l'ISO. Entreprise fictive : aucun lien avec une organisation réelle."],
    ]:
        lisez.append(l)
    lisez.column_dimensions["A"].width = 150
    lisez["A1"].font = Font(bold=True, size=13)

    feuille("Socle", ["Règle", "Section", "Intitulé", "État", "Écart", "Mesures", "Annexe A", "Source"],
            [[r["id"], r["section"], r["titre"], r["etat"], r["ecart"], ", ".join(m["id"] for m in et.mesures if r["id"] in m["regles"]),
              ", ".join(r["iso"]), r["source"]] for r in et.socle], {"C": 60, "E": 60, "B": 26})

    feuille("Événements redoutés", ["Réf.", "Valeur métier", "Événement redouté", "Gravité", "Source"],
            [[x["id"], x["vm"], x["libelle"], x["gravite"], x["source"]] for x in et.e["evenements_redoutes"]], {"C": 90})

    feuille("Couples SR-OV", ["Réf.", "Source de risque", "Objectif visé", "Motivation", "Ressources", "Activité", "Pertinence", "Retenu", "ER", "Justification", "Source"],
            [[s["id"], s["sr"], s["ov"], s["motivation"], s["ressources"], s["activite"], s["pertinence"], "oui" if s["retenu"] else "non", s["er"], s["justification"], s["source"]] for s in et.e["couples_sr_ov"]],
            {"C": 60, "J": 60})

    pp = feuille("Parties prenantes", ["Réf.", "Nom", "Catégorie", "Dépendance", "Pénétration", "Maturité", "Confiance", "Dangerosité", "Critique", "Pénétration après", "Maturité après", "Dangerosité après", "Source"],
                 [[p["id"], p["nom"], p["categorie"], p["dependance"], p["penetration"], p["maturite"], p["confiance"], f"=ROUND(D{n}*E{n}/(F{n}*G{n}),2)",
                   "oui" if p["critique"] else "non", p.get("apres", {}).get("penetration", p["penetration"]), p.get("apres", {}).get("maturite", p["maturite"]),
                   f"=ROUND(D{n}*J{n}/(K{n}*G{n}),2)", p["source"]] for n, p in enumerate(et.e["parties_prenantes"], 2)], {"B": 34, "M": 40})

    lignes = []
    for n, r in enumerate(et.e["risques"], 2):
        lignes.append([r["id"], r["libelle"], r["chemin"], r["er"], r["gravite"], r["vraisemblance"], f"=E{n}*F{n}", formule_niveau(f"G{n}", et),
                       r["residuel"]["gravite"], r["residuel"]["vraisemblance"], f"=I{n}*J{n}", formule_niveau(f"K{n}", et),
                       ", ".join(m["id"] for m in et.mesures_du_risque(r["id"])), r["justification_residuel"], r["source"]])
    ws = feuille("Risques", ["Réf.", "Scénario de risque", "Chemin", "ER", "G", "V", "Score", "Niveau", "G rés.", "V rés.", "Score rés.", "Niveau rés.", "Mesures", "Justification du résiduel", "Source"],
                 lignes, {"B": 60, "M": 40, "N": 70})
    n = len(et.e["risques"]) + 1
    couleurs = {"Faible": "C6EFCE", "Moyen": "FFEB9C", "Élevé": "F8CBAD"}
    for col in ("H", "L"):
        for nom, coul in couleurs.items():
            ws.conditional_formatting.add(f"{col}2:{col}{n}", CellIsRule(operator="equal", formula=[f'"{nom}"'], fill=PatternFill("solid", bgColor=coul)))

    feuille("Modes opératoires", ["Risque", "Mode", "Description", "Action", "Vraisemblance", "Règles du socle"],
            [[rid, i, m["nom"], a["t"], a["v"], ", ".join(str(x) for x in a["socle"])]
             for rid, blob in et.e["operationnels"].items() for i, m in enumerate(blob["modes"], 1) for a in m["actions"]], {"C": 45, "D": 60})

    feuille("Plan de traitement", ["Réf.", "Mesure", "Axe", "Risques", "Responsable", "Freins", "Coût", "Charge", "Échéance (mois après T0)", "Priorité", "Priorité guide", "Statut", "Règles socle", "Annexe A", "Source"],
            [[m["id"], m["titre"], m["groupe"], ", ".join(m["risques"] or m.get("er", [])), m["responsable"], m["freins"], m["cout"], m["charge"],
              m["echeance_mois"], m["priorite"], m.get("priorite_anssi", ""), m["statut"], ", ".join(str(r) for r in m["regles"]), ", ".join(m["iso"]), m["source"]]
             for m in et.mesures], {"B": 70, "F": 30, "N": 24, "O": 34})

    feuille("Déclaration d'applicabilité", ["N°", "Mesure", "Thème", "Applicable", "Statut", "Base du choix", "Justification", "Mesures du plan", "Risques", "Exclusion approuvée par"],
            [[cid, a["titre"], a["theme"], "Non" if soa[cid]["statut"] == "excluded" else "Oui", LIBELLE_STATUT[soa[cid]["statut"]], ", ".join(soa[cid]["base"]),
              soa[cid]["pourquoi"], ", ".join(sorted(soa[cid]["mesures"])), ", ".join(sorted(soa[cid]["risques"])), soa[cid]["approuve_par"]] for cid, a in et.annex.items()],
            {"B": 45, "G": 80})

    feuille("Indicateurs", ["Réf.", "Indicateur", "Formule", "Situation de départ", "Cible", "Fréquence", "Responsable", "Lien"],
            [[i["id"], i["nom"], i["formule"], i.get("depart", ""), i["cible"], i["frequence"], i["responsable"], i.get("lien", "")] for i in et.indicateurs],
            {"B": 50, "C": 60, "D": 30, "E": 30})
    ecrire_si_different(wb, OUT_XLSX)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="échoue si les tableaux générés sont périmés")
    args = ap.parse_args()
    et = Etude()
    errs, warns = verifier(et)
    for w in warns:
        print("AVERTISSEMENT", w)
    for e in errs:
        print("ERREUR       ", e)
    print(f"\n{len(errs)} erreur(s), {len(warns)} avertissement(s)")
    if errs:
        return 1
    soa = et.soa()
    docs = tableaux(et, soa)
    if args.check:
        perimes = [n for n, t in docs.items() if not (OUT_MD / n).is_file() or (OUT_MD / n).read_text(encoding="utf-8") != t]
        for n in perimes:
            print(f"périmé : docs/tableaux/{n} (relancer python tools/build.py)")
        pr, nouveau = readme_avec_bloc(et, soa)
        if nouveau != pr.read_text(encoding="utf-8"):
            perimes.append("README.md (bloc de synthèse)")
            print("périmé : bloc de synthèse du README (relancer python tools/build.py)")
        return 1 if perimes else 0
    OUT_MD.mkdir(parents=True, exist_ok=True)
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    for n, t in docs.items():
        (OUT_MD / n).write_text(t, encoding="utf-8", newline="\n")
        print("écrit docs/tableaux/" + n)
    pr, nouveau = readme_avec_bloc(et, soa)
    pr.write_text(nouveau, encoding="utf-8", newline="\n")
    excel(et, soa)
    print("écrit reports/" + OUT_XLSX.name)
    print("OK : données cohérentes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
